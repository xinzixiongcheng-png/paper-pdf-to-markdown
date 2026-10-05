"""PowerPoint conversion, titles, notes, charts, and image hooks."""

import base64
import inspect
import io
import os
from pathlib import Path
from typing import Any, BinaryIO, Callable, Optional
from unittest.mock import MagicMock, Mock

import pptx
import pytest
from bs4 import BeautifulSoup
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from markitdown import (
    FileConversionException,
    MarkItDown,
    MissingDependencyException,
    StreamInfo,
)
from markitdown.converters import HtmlConverter, PptxConverter, _pptx_converter


# Slide titles

BODY_TEXT = "Some body text on the slide."


def _build_pptx_with_title(title: str | None) -> io.BytesIO:
    """Build a one-slide "Title and Content" deck, optionally filling the title.

    Passing None leaves the title placeholder as PowerPoint creates it: present
    on the slide, showing "Click to add title", and carrying no text.
    """
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    if title is not None:
        slide.shapes.title.text = title
    slide.placeholders[1].text = BODY_TEXT
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def _convert_title_slide(title: str | None) -> str:
    return (
        MarkItDown()
        .convert_stream(
            _build_pptx_with_title(title), stream_info=StreamInfo(extension=".pptx")
        )
        .markdown
    )


def _heading_lines(markdown: str) -> list[str]:
    return [line for line in markdown.splitlines() if line.startswith("#")]


def test_untouched_title_placeholder_produces_no_heading() -> None:
    markdown = _convert_title_slide(None)

    assert BODY_TEXT in markdown
    assert _heading_lines(markdown) == []


def test_empty_title_produces_no_heading() -> None:
    markdown = _convert_title_slide("")

    assert BODY_TEXT in markdown
    assert _heading_lines(markdown) == []


def test_whitespace_only_title_produces_no_heading() -> None:
    markdown = _convert_title_slide("   ")

    assert BODY_TEXT in markdown
    assert _heading_lines(markdown) == []


def test_title_with_text_is_still_emitted() -> None:
    markdown = _convert_title_slide("Quarterly Results")

    assert "# Quarterly Results" in markdown
    assert BODY_TEXT in markdown


# Speaker notes


def _build_pptx_with_optional_notes(notes: str | None) -> io.BytesIO:
    """Build a one-slide deck in memory, optionally attaching a notes slide."""
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Slide title"
    if notes is not None:
        # Touching notes_slide creates the part, which is what PowerPoint does
        # for a deck whose notes pane has been opened.
        slide.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def _convert_notes_slide(notes: str | None) -> str:
    return (
        MarkItDown()
        .convert_stream(
            _build_pptx_with_optional_notes(notes),
            stream_info=StreamInfo(extension=".pptx"),
        )
        .markdown
    )


def test_empty_notes_slide_produces_no_notes_heading() -> None:
    markdown = _convert_notes_slide("")

    assert "Slide title" in markdown
    assert "### Notes:" not in markdown


def test_whitespace_only_notes_produce_no_notes_heading() -> None:
    markdown = _convert_notes_slide("   \n\n  ")

    assert "Slide title" in markdown
    assert "### Notes:" not in markdown


def test_slide_without_a_notes_slide_produces_no_notes_heading() -> None:
    markdown = _convert_notes_slide(None)

    assert "Slide title" in markdown
    assert "### Notes:" not in markdown


def test_notes_with_text_are_still_emitted() -> None:
    markdown = _convert_notes_slide("Remember to mention the budget.")

    assert "### Notes:\nRemember to mention the budget." in markdown


# Missing text

# Empty XML text nodes should not drop surrounding slide content. python-pptx
# normalizes <a:t/> to ""; it does not return None for shape/text-frame text.

TEST_FILES_DIR = os.path.join(os.path.dirname(__file__), "test_files")

PPTX_FIXTURE = os.path.join(TEST_FILES_DIR, "test.pptx")


@pytest.mark.parametrize("empty_shape", ["title", "body"])
def test_pptx_fixture_contains_empty_text_run(empty_shape: str) -> None:
    path = Path(TEST_FILES_DIR) / f"pptx_empty_{empty_shape}_run.pptx"
    # Confirm the fixture contains an empty text node, not just no runs.
    reloaded = Presentation(path).slides[-1]
    shape = (
        reloaded.shapes.title if empty_shape == "title" else reloaded.placeholders[1]
    )
    assert shape.text_frame.paragraphs[0].runs[0]._r.t.text is None
    assert shape.text == ""


@pytest.mark.parametrize("notes_state", ["empty-run", "missing-body-frame"])
def test_pptx_empty_notes_have_no_heading(notes_state: str) -> None:
    suffix = notes_state.replace("-", "_")
    path = Path(TEST_FILES_DIR) / f"pptx_notes_{suffix}.pptx"
    reloaded = Presentation(path).slides[-1]
    assert reloaded.has_notes_slide
    frame = reloaded.notes_slide.notes_text_frame
    if notes_state == "empty-run":
        assert frame.paragraphs[0].runs[0]._r.t.text is None
        assert frame.text == ""
    else:
        assert frame is None
    markdown = MarkItDown().convert(path).markdown

    # Earlier slides retain their notes; only the appended slide has empty notes.
    last_slide = markdown.rsplit("<!-- Slide number:", 1)[1]
    assert "### Notes:" not in last_slide


# SVG pictures

# Tests for PPTX SVG images that lack a rasterized fallback.
#
# PowerPoint stores an SVG picture as an ``<a:blip>`` whose ``r:embed`` points to
# a rasterized PNG fallback, plus an ``<asvg:svgBlip>`` extension that points to
# the SVG. When a picture has no raster fallback the ``<a:blip>`` has no
# ``r:embed`` at all, so python-pptx's ``shape.image`` raises
# ``ValueError("no embedded image")``. The converter must handle this gracefully
# (resolving the SVG blip directly) instead of failing the whole conversion.

# A tiny synthetic PPTX whose only picture is an SVG without a rasterized
# fallback: the <a:blip> has no r:embed, only an <asvg:svgBlip> extension
# pointing to an embedded SVG. Its alt text is "Red square SVG".
SVG_NO_FALLBACK_PPTX = os.path.join(TEST_FILES_DIR, "test_svg_no_fallback.pptx")

_SVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"


def _first_picture_shape(pptx_path):
    presentation = Presentation(pptx_path)
    converter = PptxConverter()
    for slide in presentation.slides:
        for shape in slide.shapes:
            if converter._is_picture(shape):
                return converter, shape
    raise AssertionError(f"No picture shape found in {pptx_path}")


def test_pptx_svg_without_raster_fallback() -> None:
    md = MarkItDown()

    # Default conversion should not raise and should emit the image alt text.
    result = md.convert(SVG_NO_FALLBACK_PPTX)
    assert "Red square SVG" in result.markdown

    # keep_data_uris used to crash with ValueError("no embedded image"). It
    # should now embed the SVG as a data URI.
    result = md.convert(SVG_NO_FALLBACK_PPTX, keep_data_uris=True)
    assert "data:image/svg+xml;base64," in result.markdown


def test_get_image_info_resolves_svg_blip_without_fallback() -> None:
    # Unit-level check: _get_image_info must resolve the raw SVG blob directly
    # from the <asvg:svgBlip> extension when shape.image raises because there is
    # no rasterized fallback (no r:embed on the <a:blip>).
    converter, shape = _first_picture_shape(SVG_NO_FALLBACK_PPTX)

    blob, content_type, _filename = converter._get_image_info(shape)

    assert blob is not None and len(blob) > 0
    assert content_type == "image/svg+xml"
    assert b"<svg" in blob[:512].lower()


class _FakePart:
    """Minimal part whose related_part always resolves to a truthy object."""

    def related_part(self, rid):
        return object()


class _FakeSvgPlaceholderShape:
    """A placeholder shape whose ``image`` raises like an SVG-only placeholder.

    ``_is_picture`` used to rely on ``hasattr(shape, "image")`` which only
    swallows ``AttributeError`` and let this ``ValueError`` propagate, failing
    even the default conversion.
    """

    shape_type = MSO_SHAPE_TYPE.PLACEHOLDER

    def __init__(self):
        xml = (
            '<p:pic xmlns:p="http://schemas.openxmlformats.org/'
            'presentationml/2006/main" xmlns:a="http://schemas.'
            'openxmlformats.org/drawingml/2006/main" xmlns:r="http://'
            'schemas.openxmlformats.org/officeDocument/2006/relationships" '
            'xmlns:asvg="%s"><a:blip><a:extLst><a:ext uri="{96DAC541-7B7A-'
            '43D3-8B79-37D633B846F1}"><asvg:svgBlip r:embed="rId9"/></a:ext>'
            "</a:extLst></a:blip></p:pic>" % _SVG_NS
        )
        self._element = etree.fromstring(xml)
        self.part = _FakePart()

    @property
    def image(self):
        raise ValueError("no embedded image")


def test_is_picture_true_for_svg_placeholder() -> None:
    # _is_picture must not let shape.image's ValueError propagate for SVG
    # placeholders lacking a raster fallback; it should still report a picture
    # because an embedded SVG blip is present.
    converter = PptxConverter()
    assert converter._is_picture(_FakeSvgPlaceholderShape()) is True


# Image hooks

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4z8AAAAM"
    "BAQDJ/pLvAAAAAElFTkSuQmCC"
)

_GIF = base64.b64decode("R0lGODlhAQABAIAAAAAAAP///yH5BAEAAAAALAAAAAABAAEAAAIBTAA7")

_INFO = StreamInfo(extension=".pptx")

_FILES = Path(__file__).parent / "test_files"


def _presentation(
    images: tuple[bytes, ...] = (_PNG,), *, native_content: bool = False
) -> io.BytesIO:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Deck title"
    slide.shapes.title.top = 0
    # Insert in reverse order to distinguish shape order from reading order.
    for index in reversed(range(len(images))):
        image = slide.shapes.add_picture(
            io.BytesIO(images[index]), 0, Inches(index + 1), width=Inches(0.5)
        )
        image.name = f"Picture {index + 1}"
        image._element._nvXxPr.cNvPr.set("descr", f"Alt [{index + 1}]\r\n& text")
    if native_content:
        table = slide.shapes.add_table(2, 2, 0, Inches(4), Inches(4), Inches(1)).table
        for cell, value in zip(
            (table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)),
            ("Item", "Details", "A", "B & <literal>"),
        ):
            cell.text = value
        chart_data = CategoryChartData()
        chart_data.categories = ["Q1", "Q2"]
        chart_data.add_series("Revenue", (10, 20))
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            0,
            Inches(5),
            Inches(4),
            Inches(1),
            chart_data,
        ).chart
        chart.has_title = True
        chart.chart_title.text_frame.text = "Sales"
        slide.notes_slide.notes_text_frame.text = "Speaker notes"
        next_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        next_slide.shapes.title.text = ""
        next_slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Closing"
        next_slide.notes_slide.notes_text_frame.text = " \n "
    stream = io.BytesIO()
    presentation.save(stream)
    stream.seek(0)
    return stream


class _ImageConverter(PptxConverter):
    def __init__(self, render: Callable[..., Optional[str]]):
        super().__init__()
        self.render = render

    def _image_to_html(
        self, image_stream: BinaryIO, stream_info: StreamInfo, **kwargs: Any
    ) -> Optional[str]:
        return self.render(image_stream, stream_info, **kwargs)


def test_hook_does_not_change_public_signature() -> None:
    assert list(inspect.signature(PptxConverter.convert).parameters) == [
        "self",
        "file_stream",
        "stream_info",
        "kwargs",
    ]
    assert list(inspect.signature(PptxConverter.__init__).parameters) == ["self"]
    assert _ImageConverter.convert is PptxConverter.convert


@pytest.mark.parametrize("via_dispatcher", [False, True])
def test_inherited_hook_receives_real_images_and_options_in_reading_order(
    via_dispatcher: bool,
) -> None:
    seen = []
    streams = []
    service = object()

    def render(stream: BinaryIO, info: StreamInfo, **kwargs: Any) -> str:
        assert stream.tell() == 0 and stream.seekable()
        data = stream.read()
        stream.seek(0)
        assert stream.read() == data
        seen.append((data, info, kwargs))
        streams.append(stream)
        return f"<strong>Image {len(seen)}</strong>"

    class InheritedImages(_ImageConverter):
        pass

    converter = InheritedImages(render)
    source = _presentation((_GIF, _PNG, _GIF))
    original = source.getvalue()
    info = StreamInfo(
        extension=".pptx",
        filename="deck.pptx",
        url="https://example.test/deck.pptx",
        local_path="/deck.pptx",
    )
    options: dict[str, Any] = {"ocr_service": service, "escape_underscores": False}
    if via_dispatcher:
        md = MarkItDown()
        md.register_converter(converter, priority=-1)
        result = md.convert_stream(source, stream_info=info, **options)
    else:
        source.seek(7)
        result = converter.convert(source, info, **options)

    assert [entry[0] for entry in seen] == [_GIF, _PNG, _GIF]
    assert [entry[1] for entry in seen] == [
        StreamInfo(
            mimetype=f"image/{ext}", extension=f".{ext}", filename=f"image.{ext}"
        )
        for ext in ("gif", "png", "gif")
    ]
    assert all(entry[2]["ocr_service"] is service for entry in seen)
    assert all(entry[2]["escape_underscores"] is False for entry in seen)
    assert result.markdown == (
        "<!-- Slide number: 1 -->\n# Deck title\n"
        "\n**Image 1**\n\n**Image 2**\n\n**Image 3**"
    )
    assert result.title is None
    assert all(stream.closed for stream in streams)
    assert not source.closed and source.getvalue() == original


@pytest.mark.parametrize(
    ("fallback", "keep_data_uris"), [(None, False), ("", True), (" \r\n\t", False)]
)
def test_declining_hook_preserves_native_output_and_metadata(
    fallback: Optional[str], keep_data_uris: bool
) -> None:
    expected = PptxConverter().convert(
        _presentation(native_content=True), _INFO, keep_data_uris=keep_data_uris
    )
    render = Mock(return_value=fallback)
    actual = _ImageConverter(render).convert(
        _presentation(native_content=True), _INFO, keep_data_uris=keep_data_uris
    )

    assert actual.markdown == expected.markdown
    assert actual.title == expected.title
    render.assert_called_once()


@pytest.mark.parametrize("keep_data_uris", [False, True])
def test_declining_hook_preserves_real_presentation(keep_data_uris: bool) -> None:
    data = (_FILES / "test.pptx").read_bytes()
    expected = PptxConverter().convert(
        io.BytesIO(data), _INFO, keep_data_uris=keep_data_uris
    )
    render = Mock(return_value=None)

    actual = _ImageConverter(render).convert(
        io.BytesIO(data), _INFO, keep_data_uris=keep_data_uris
    )

    assert actual.markdown == expected.markdown
    assert actual.title == expected.title
    assert render.call_count > 0


def test_native_converter_does_not_add_image_html_processing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class NativeImages(PptxConverter):
        pass

    for converter in (PptxConverter(), NativeImages()):
        convert_html = Mock(side_effect=AssertionError("unexpected HTML"))
        monkeypatch.setattr(converter._html_converter, "convert_string", convert_html)
        assert (
            "![Alt 1 & text](Picture1.jpg)"
            in converter.convert(_presentation(), _INFO).markdown
        )
        convert_html.assert_not_called()


def test_override_can_delegate_to_super() -> None:
    class NativeImages(PptxConverter):
        def _image_to_html(self, image_stream, stream_info, **kwargs):
            return super()._image_to_html(image_stream, stream_info, **kwargs)

    assert (
        NativeImages().convert(_presentation(), _INFO).markdown
        == PptxConverter().convert(_presentation(), _INFO).markdown
    )


def test_custom_fragment_uses_html_options_without_reprocessing_generated_images(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fragment = (
        "<h2>Image heading</h2><p>A_B &amp; &lt;literal&gt;<br>second</p>"
        '<img src="generated.png" alt="generated"><script>discard</script>'
    )
    render = Mock(return_value=fragment)
    converter = _ImageConverter(render)
    convert_html = Mock(wraps=converter._html_converter.convert_string)
    monkeypatch.setattr(converter._html_converter, "convert_string", convert_html)
    options: dict[str, Any] = {
        "escape_underscores": False,
        "heading_style": "underlined",
    }

    result = converter.convert(_presentation(), _INFO, **options)

    convert_html.assert_called_once_with(
        str(BeautifulSoup(fragment, "html.parser")), **options
    )
    expected = HtmlConverter().convert_string(fragment, **options).markdown
    assert result.markdown == "<!-- Slide number: 1 -->\n# Deck title\n\n" + expected
    render.assert_called_once()


def test_replacements_keep_native_table_chart_notes_and_slide_placement() -> None:
    render = Mock(side_effect=["<p>first</p>", None, "<p>third</p>"])
    result = _ImageConverter(render).convert(
        _presentation((_PNG, _GIF, _PNG), native_content=True), _INFO
    )
    markdown = result.markdown
    ordered = [
        "# Deck title",
        "first",
        "![Alt 2 & text](Picture2.jpg)",
        "third",
        "| Item | Details |",
        "| A | B & <literal> |",
        "### Chart: Sales",
        "| Q1 | 10.0 |",
        "### Notes:\nSpeaker notes",
        "<!-- Slide number: 2 -->",
        "Closing",
    ]
    assert [markdown.index(text) for text in ordered] == sorted(
        markdown.index(text) for text in ordered
    )
    assert markdown.count("### Notes:") == 1
    assert "\n# \n" not in markdown


def test_group_images_follow_native_negative_and_zero_coordinate_order() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_picture(io.BytesIO(_PNG), 0, 0, width=Inches(0.5))
    group.shapes.add_picture(io.BytesIO(_GIF), 0, -Inches(1), width=Inches(0.5))
    stream = io.BytesIO()
    presentation.save(stream)
    stream.seek(0)
    seen = []

    def render(image_stream, info, **kwargs):
        seen.append(image_stream.read())
        return f"<p>image{len(seen)}</p>"

    result = _ImageConverter(render).convert(stream, _INFO)

    assert seen == [_GIF, _PNG]
    assert result.markdown == "<!-- Slide number: 1 -->\n\nimage1\n\nimage2"


def test_picture_placeholder_invokes_hook() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[8])
    placeholder = next(
        shape for shape in slide.placeholders if hasattr(shape, "insert_picture")
    )
    placeholder.insert_picture(io.BytesIO(_PNG))
    stream = io.BytesIO()
    presentation.save(stream)
    stream.seek(0)
    render = Mock(return_value="<p>placeholder</p>")

    result = _ImageConverter(render).convert(stream, _INFO)

    assert result.markdown == "<!-- Slide number: 1 -->\n\nplaceholder"
    render.assert_called_once()


def test_svg_without_raster_fallback_reaches_image_hook() -> None:
    seen = []

    def render(stream, info, **kwargs):
        seen.append((stream.read(), info))
        return "<p>SVG content</p>"

    with (_FILES / "test_svg_no_fallback.pptx").open("rb") as stream:
        result = _ImageConverter(render).convert(stream, _INFO)

    assert len(seen) == 1 and b"<svg" in seen[0][0]
    assert seen[0][1].mimetype == "image/svg+xml"
    assert seen[0][1].extension == ".svg"
    assert seen[0][1].url is None
    assert "SVG content" in result.markdown


def test_svg_picture_placeholder_uses_native_resolution_and_image_hook() -> None:
    presentation = Presentation(str(_FILES / "test_svg_no_fallback.pptx"))
    picture = next(
        shape
        for slide in presentation.slides
        for shape in slide.shapes
        if PptxConverter()._is_picture(shape)
    )
    nonvisual = picture._element.xpath("./p:nvPicPr/p:nvPr")[0]
    etree.SubElement(
        nonvisual,
        "{http://schemas.openxmlformats.org/presentationml/2006/main}ph",
        type="pic",
        idx="1",
    )
    stream = io.BytesIO()
    presentation.save(stream)
    stream.seek(0)
    presentation = Presentation(stream)
    placeholder = next(
        shape
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER
    )
    with pytest.raises(ValueError, match="no embedded image"):
        _ = placeholder.image
    stream.seek(0)
    render = Mock(return_value="<p>SVG placeholder</p>")

    result = _ImageConverter(render).convert(stream, _INFO)

    assert "SVG placeholder" in result.markdown
    assert render.call_args.args[1].mimetype == "image/svg+xml"
    render.assert_called_once()


def test_missing_image_bytes_keep_native_placeholder(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    render = Mock(side_effect=AssertionError("unexpected image"))
    converter = _ImageConverter(render)
    monkeypatch.setattr(
        converter, "_get_image_info", Mock(return_value=(None, None, None))
    )

    result = converter.convert(_presentation(), _INFO, keep_data_uris=True)

    assert "![Alt 1 & text](Picture1.jpg)" in result.markdown
    render.assert_not_called()


@pytest.mark.parametrize("value", [False, 123, b"<p>not text</p>"])
def test_non_string_fragments_fail_explicitly(value: Any) -> None:
    with pytest.raises(TypeError, match="HTML string or None"):
        _ImageConverter(Mock(return_value=value)).convert(_presentation(), _INFO)


@pytest.mark.parametrize(
    "fragment",
    [
        "<html><body>document</body></html>",
        "<head><title>document</title></head>",
        "<body>document</body>",
        "<!DOCTYPE html><p>document</p>",
    ],
)
def test_full_documents_fail_explicitly(fragment: str) -> None:
    with pytest.raises(ValueError, match="fragment, not a document"):
        _ImageConverter(Mock(return_value=fragment)).convert(_presentation(), _INFO)


def test_hook_errors_propagate_close_stream_and_use_dispatcher_fallback() -> None:
    streams = []
    error = RuntimeError("image renderer failed")

    def render(stream, info, **kwargs):
        streams.append(stream)
        raise error

    converter = _ImageConverter(render)
    with pytest.raises(RuntimeError) as caught:
        converter.convert(_presentation(), _INFO)
    assert caught.value is error
    without_fallback = MarkItDown(enable_builtins=False)
    without_fallback.register_converter(converter)
    with pytest.raises(FileConversionException) as aggregate:
        without_fallback.convert_stream(_presentation(), stream_info=_INFO)
    assert aggregate.value.attempts is not None
    assert any(
        attempt.exc_info and attempt.exc_info[1] is error
        for attempt in aggregate.value.attempts
    )
    md = MarkItDown()
    md.register_converter(converter, priority=-1)
    assert (
        md.convert_stream(_presentation(), stream_info=_INFO).markdown
        == PptxConverter().convert(_presentation(), _INFO).markdown
    )
    assert streams and all(stream.closed for stream in streams)


@pytest.mark.parametrize("caption_result", [None, "", RuntimeError("caption failed")])
def test_caption_failure_precedes_hook_with_fresh_image_stream(
    monkeypatch: pytest.MonkeyPatch, caption_result: Any
) -> None:
    calls = []
    client = object()

    def caption(stream, info, **kwargs):
        assert stream.tell() == 0 and stream.read() == _PNG
        assert info.mimetype == "image/png" and info.extension == ".png"
        assert kwargs == {"client": client, "model": "vision", "prompt": "Describe"}
        calls.append("caption")
        if isinstance(caption_result, Exception):
            raise caption_result
        return caption_result

    def render(stream, info, **kwargs):
        assert stream.tell() == 0 and stream.read() == _PNG
        calls.append("hook")
        return "<p>recognized</p>"

    monkeypatch.setattr(_pptx_converter, "llm_caption", caption)
    result = _ImageConverter(render).convert(
        _presentation(),
        _INFO,
        llm_client=client,
        llm_model="vision",
        llm_prompt="Describe",
    )

    assert calls == ["caption", "hook"]
    assert "recognized" in result.markdown


def test_successful_caption_keeps_native_markdown_without_hook_or_html(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    caption = Mock(return_value="**Caption** [label]\nwith_under & <angle>")
    monkeypatch.setattr(_pptx_converter, "llm_caption", caption)
    render = Mock(side_effect=AssertionError("caption must take precedence"))
    converter = _ImageConverter(render)
    convert_html = Mock(side_effect=AssertionError("caption is not image HTML"))
    monkeypatch.setattr(converter._html_converter, "convert_string", convert_html)

    result = converter.convert(
        _presentation(), _INFO, llm_client=object(), llm_model="vision"
    )

    assert result.markdown == (
        "<!-- Slide number: 1 -->\n# Deck title\n\n"
        "![**Caption** label with_under & <angle> Alt 1 & text](Picture1.jpg)"
    )
    caption.assert_called_once()
    render.assert_not_called()
    convert_html.assert_not_called()


def test_missing_optional_dependencies_fail_before_hook(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    error = ImportError("python-pptx is unavailable")
    monkeypatch.setattr(
        _pptx_converter, "_dependency_exc_info", (ImportError, error, None)
    )
    render = Mock()
    with pytest.raises(MissingDependencyException) as caught:
        _ImageConverter(render).convert(io.BytesIO(b""), _INFO)
    assert caught.value.__cause__ is error
    render.assert_not_called()


# Conversion regressions


def test_pptx_chart_multi_series_conversion() -> None:
    """Charts with multiple series and many categories must convert correctly.

    Regression test for the slow path in PptxConverter._convert_chart_to_markdown,
    where ``series.values[idx]`` was evaluated inside the (category x series) loop.
    In python-pptx each ``series.values`` access rescans the cached points via
    XPath (O(n) per lookup), so the old code was O(n^2) per series (and rebuilt
    the whole tuple for every category), making large charts extremely slow.

    The values are now materialized once per series. This test builds a chart
    with enough categories that the regressed code path would be pathologically
    slow, and verifies the resulting Markdown table is correct across series.
    """
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    n_categories = 200
    categories = [f"C{i}" for i in range(n_categories)]
    series_a = [float(i) for i in range(n_categories)]
    series_b = [float(i * 2) for i in range(n_categories)]

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series("Series A", series_a)
    chart_data.add_series("Series B", series_b)
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(8),
        Inches(5),
        chart_data,
    )

    buffer = io.BytesIO()
    presentation.save(buffer)
    buffer.seek(0)

    result = MarkItDown().convert_stream(buffer, file_extension=".pptx")
    md = result.markdown

    # Both series headers are present
    assert "Series A" in md
    assert "Series B" in md
    # First and last categories are present (nothing truncated)
    assert "| C0 |" in md
    assert f"| C{n_categories - 1} |" in md
    # A representative row carries the correct value for each series
    assert "| C10 | 10.0 | 20.0 |" in md


def test_pptx_converter_treats_none_llm_caption_as_empty(monkeypatch) -> None:
    from markitdown.converters import _pptx_converter

    calls = 0

    def none_caption(*args, **kwargs):
        nonlocal calls
        calls += 1
        return None

    monkeypatch.setattr(_pptx_converter, "llm_caption", none_caption)

    result = MarkItDown().convert(
        os.path.join(TEST_FILES_DIR, "test.pptx"),
        llm_client=MagicMock(),
        llm_model="test-model",
    )

    assert calls > 0
    assert (
        "![This phrase of the caption is Human-written.](Picture4.jpg)"
        in result.markdown
    )


def _chart_presentation(title: str | None) -> io.BytesIO:
    # Modify an in-memory copy of an existing deck; keep its fixture intact.
    presentation = Presentation(Path(PPTX_FIXTURE))
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["Cat 1"]
    data.add_series("Series 1", [10.0])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(5),
        Inches(3),
        data,
    ).chart
    chart.has_title = True
    if title is not None:
        chart.chart_title.text_frame.text = title
    assert chart.chart_title.has_text_frame is (title is not None)
    stream = io.BytesIO()
    presentation.save(stream)
    stream.seek(0)
    return stream


@pytest.mark.parametrize("title", [None, "Revenue"])
def test_pptx_chart_title_text_frame(title: str | None) -> None:
    stream = _chart_presentation(title)
    # Reload the serialized package so the real parser reads title XML.
    chart = Presentation(stream).slides[-1].shapes[-1].chart
    converter = PptxConverter()
    result = converter._convert_chart_to_markdown(chart)
    heading = "### Chart" if title is None else "### Chart: Revenue"
    assert result.strip().splitlines()[0] == heading
    assert "Cat 1" in result
    assert "Series 1" in result
    assert "10.0" in result
    # Accessing a title without a text frame must not create one.
    assert chart.chart_title.has_text_frame is (title is not None)
    stream.seek(0)
    markdown = converter.convert(stream, StreamInfo(extension=".pptx")).markdown
    assert result.strip() in markdown


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__]))
