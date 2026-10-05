"""Spreadsheet conversion, compatibility repairs, and image hooks."""

import inspect
import io
import zipfile
from pathlib import Path
from typing import Any, BinaryIO, Callable, Optional
from unittest.mock import Mock
from xml.etree import ElementTree as ET

import openpyxl
import pytest
from PIL import Image
from bs4 import BeautifulSoup
from defusedxml.common import EntitiesForbidden
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.drawing.spreadsheet_drawing import (
    AbsoluteAnchor,
    AnchorMarker,
    OneCellAnchor,
    TwoCellAnchor,
)

from markitdown import MarkItDown, MissingDependencyException, StreamInfo
from markitdown.converter_utils import _xlsx_images
from markitdown.converters import XlsConverter, XlsxConverter, _xlsx_converter


# Image hooks

_INFO = StreamInfo(extension=".xlsx")


def _png(color: str) -> bytes:
    stream = io.BytesIO()
    Image.new("RGB", (2, 2), color).save(stream, "PNG")
    return stream.getvalue()


_RED, _BLUE = _png("red"), _png("blue")


def _workbook(*, images: bool = True) -> bytes:
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "Second"
    first.append(["Header_one", None])
    first.append(["<native> & *value*", 12])
    last = workbook.create_sheet("First")
    last.append(["Other"])
    last.append(["End"])
    if images:
        for data, anchor in (
            (_RED, OneCellAnchor(_from=AnchorMarker(col=26, row=2))),
            (
                _BLUE,
                TwoCellAnchor(
                    _from=AnchorMarker(col=2, row=4), to=AnchorMarker(col=4, row=7)
                ),
            ),
            (_RED, AbsoluteAnchor()),
        ):
            image = SheetImage(io.BytesIO(data))
            image.anchor = anchor
            first.add_image(image)
        last.add_image(SheetImage(io.BytesIO(_BLUE)), "Z10")
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _rewrite(data: bytes, change: Callable[[str, bytes], bytes]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        with zipfile.ZipFile(output, "w") as target:
            for entry in source.infolist():
                target.writestr(entry, change(entry.filename, source.read(entry)))
    return output.getvalue()


class _ImageConverter(XlsxConverter):
    def __init__(self, render: Callable[..., Optional[str]]):
        super().__init__()
        self.render = render

    def _image_to_html(
        self, image_stream: BinaryIO, stream_info: StreamInfo, **kwargs: Any
    ) -> Optional[str]:
        return self.render(image_stream, stream_info, **kwargs)


def _convert(converter: XlsxConverter, data: bytes, **kwargs: Any) -> str:
    return converter.convert(io.BytesIO(data), _INFO, **kwargs).markdown


def test_image_hook_keeps_native_public_signature_and_acceptance() -> None:
    assert list(inspect.signature(XlsxConverter.convert).parameters) == [
        "self",
        "file_stream",
        "stream_info",
        "kwargs",
    ]
    assert list(inspect.signature(XlsxConverter.__init__).parameters) == ["self"]
    assert _ImageConverter.convert is XlsxConverter.convert
    converter = XlsxConverter()
    assert converter.accepts(io.BytesIO(), StreamInfo(extension=".XLSX"))
    assert converter.accepts(
        io.BytesIO(),
        StreamInfo(
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    assert not converter.accepts(io.BytesIO(), StreamInfo(extension=".xls"))


def test_hook_receives_original_images_metadata_options_and_legacy_anchor_order() -> (
    None
):
    seen = []
    streams = []
    service = object()

    def render(stream: BinaryIO, info: StreamInfo, **kwargs: Any) -> str:
        assert stream.seekable() and stream.tell() == 0
        image = stream.read()
        stream.seek(0)
        assert stream.read() == image
        seen.append((image, info, kwargs))
        streams.append(stream)
        return f"<p><strong>image_{len(seen)}</strong></p>"

    class InheritedImages(_ImageConverter):
        pass

    source = io.BytesIO(_workbook())
    original = source.getvalue()
    source.seek(7)
    info = StreamInfo(
        extension=".xlsx",
        filename="parent.xlsx",
        local_path="/parent.xlsx",
        url="https://example.test/parent.xlsx",
    )
    result = InheritedImages(render).convert(
        source,
        info,
        ocr_service=service,
        heading_style="atx",
        escape_underscores=False,
        filename="parent-option.xlsx",
        url="https://example.test/option.xlsx",
    )

    # Preserve openpyxl's absolute, one-cell, two-cell traversal, not ZIP order.
    assert [(data, metadata) for data, metadata, _ in seen] == [
        (
            _RED,
            StreamInfo(extension=".png", mimetype="image/png", filename="image3.png"),
        ),
        (
            _RED,
            StreamInfo(extension=".png", mimetype="image/png", filename="image1.png"),
        ),
        (
            _BLUE,
            StreamInfo(extension=".png", mimetype="image/png", filename="image2.png"),
        ),
        (
            _BLUE,
            StreamInfo(extension=".png", mimetype="image/png", filename="image4.png"),
        ),
    ]
    assert all(options["ocr_service"] is service for _, _, options in seen)
    assert all(options["filename"] == "parent-option.xlsx" for _, _, options in seen)
    expected_order = [
        "## Second",
        "| <native> & \\*value\\* | 12 |",
        "### Images in this sheet:",
        "**image_1**",
        "**image_2**",
        "**image_3**",
        "## First",
        "| End |",
        "**image_4**",
    ]
    offsets = [result.markdown.index(text) for text in expected_order]
    assert offsets == sorted(offsets)
    assert "Image at " not in result.markdown
    assert all(stream.closed for stream in streams)
    assert not source.closed and source.getvalue() == original


@pytest.mark.parametrize(
    ("template", "expected"),
    [
        ("<strong>{}</strong>", "**one**\n\n**two**\n\n**three**"),
        ("{}", "one\n\ntwo\n\nthree"),
        ("{}<br>line", "one  \nline\n\ntwo  \nline\n\nthree  \nline"),
        (
            "<p>{}</p><p>paragraph</p>",
            "one\n\nparagraph\n\ntwo\n\nparagraph\n\nthree\n\nparagraph",
        ),
    ],
)
def test_image_fragments_are_separate_valid_blocks(
    monkeypatch: pytest.MonkeyPatch, template: str, expected: str
) -> None:
    render = Mock(
        side_effect=[template.format(text) for text in ("one", "two", "three", "four")]
    )
    converter = _ImageConverter(render)
    convert_html = Mock(wraps=converter._html_converter.convert_string)
    monkeypatch.setattr(converter._html_converter, "convert_string", convert_html)

    result = _convert(converter, _workbook())

    assert f"### Images in this sheet:\n\n{expected}\n\n## First\n" in result
    soup = BeautifulSoup(convert_html.call_args_list[1].args[0], "html.parser")
    assert [tag.name for tag in soup.find_all(recursive=False)] == [
        "h3",
        "div",
        "div",
        "div",
    ]
    assert not soup.select("p div")


@pytest.mark.parametrize("fragment", [None, "", " \n\t"])
def test_declining_hook_is_exactly_native(fragment: Optional[str]) -> None:
    data = _workbook()
    options = {"escape_underscores": False, "heading_style": "underlined"}
    assert _convert(_ImageConverter(Mock(return_value=fragment)), data, **options) == (
        _convert(XlsxConverter(), data, **options)
    )


def test_default_converter_does_not_parse_images_or_change_native_output(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        _xlsx_images, "_XlsxImages", Mock(side_effect=AssertionError("extra parsing"))
    )
    reader = Mock(wraps=_xlsx_converter.pd.read_excel)
    monkeypatch.setattr(_xlsx_converter.pd, "read_excel", reader)

    class NativeSubclass(XlsxConverter):
        pass

    expected = (
        "## Second\n"
        "| Header\\_one | Unnamed: 1 |\n"
        "| --- | --- |\n"
        "| <native> & \\*value\\* | 12 |\n\n"
        "## First\n| Other |\n| --- |\n| End |"
    )
    for converter in (XlsxConverter(), NativeSubclass()):
        assert _convert(converter, _workbook()) == expected
    assert reader.call_count == 2


def test_image_hook_uses_repaired_package_and_repair_stream_is_closed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data = _rewrite(
        _workbook(),
        lambda name, content: (
            content.replace(b"<sheetView ", b'<sheetView showZeroes="0" ', 1)
            if name == "xl/worksheets/sheet1.xml"
            else content
        ),
    )
    repaired_streams = []
    original_repair = _xlsx_converter._repair_sheetview_show_zeroes

    def repair(*args: Any) -> io.BytesIO:
        stream = original_repair(*args)
        repaired_streams.append(stream)
        return stream

    original_images = _xlsx_images._XlsxImages

    def images(stream: BinaryIO):
        assert stream is repaired_streams[-1]
        with zipfile.ZipFile(stream) as archive:
            sheet = archive.read("xl/worksheets/sheet1.xml")
            assert b'showZeros="0"' in sheet and b"showZeroes" not in sheet
        return original_images(stream)

    monkeypatch.setattr(_xlsx_converter, "_repair_sheetview_show_zeroes", repair)
    monkeypatch.setattr(_xlsx_images, "_XlsxImages", images)
    render = Mock(return_value="<p>recognized</p>")
    result = _convert(_ImageConverter(render), data)
    assert result.count("recognized") == 4
    assert render.call_count == 4
    assert all(stream.closed for stream in repaired_streams)


@pytest.mark.parametrize(
    ("fragment", "error"),
    [
        (17, TypeError),
        (b"<p>bytes</p>", TypeError),
        ("<!doctype html><p>text</p>", ValueError),
        ("<html><body>text</body></html>", ValueError),
        ("<head><title>title</title></head>", ValueError),
    ],
)
def test_invalid_hook_results_are_rejected(
    fragment: Any, error: type[Exception]
) -> None:
    with pytest.raises(error, match="_image_to_html must return"):
        _convert(_ImageConverter(Mock(return_value=fragment)), _workbook())


def test_hook_failure_closes_image_stream_and_remains_visible() -> None:
    streams = []
    error = RuntimeError("image failure")

    def render(stream: BinaryIO, *args: Any, **kwargs: Any) -> str:
        streams.append(stream)
        raise error

    source = io.BytesIO(_workbook())
    with pytest.raises(RuntimeError) as caught:
        _ImageConverter(render).convert(source, _INFO)
    assert caught.value is error
    assert streams and all(stream.closed for stream in streams)
    assert not source.closed


def test_native_table_failure_is_not_silently_skipped(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    error = TypeError("not the repairable error")
    monkeypatch.setattr(_xlsx_converter.pd, "read_excel", Mock(side_effect=error))
    render = Mock(return_value="<p>recognized</p>")
    with pytest.raises(TypeError) as caught:
        _convert(_ImageConverter(render), _workbook())
    assert caught.value is error
    render.assert_not_called()


def test_package_relationships_preserve_unusual_image_names_types_and_bytes() -> None:
    data = _workbook()
    svg = b'<svg xmlns="http://www.w3.org/2000/svg"><text>vector</text></svg>'
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        with zipfile.ZipFile(output, "w") as target:
            for entry in source.infolist():
                content = source.read(entry)
                name = entry.filename
                if name == "xl/media/image1.png":
                    name, content = "xl/media/vector image.SVG", svg
                elif name.endswith(".rels"):
                    content = content.replace(
                        b"/xl/media/image1.png", b"../media/vector%20image.SVG"
                    )
                elif name == "[Content_Types].xml":
                    content = content.replace(
                        b"</Types>",
                        b'<Override PartName="/xl/media/vector%20image.SVG" '
                        b'ContentType="image/svg+xml"/></Types>',
                    )
                target.writestr(name, content)
    render = Mock(return_value="<p>recognized</p>")

    result = _convert(_ImageConverter(render), output.getvalue())

    assert result.count("recognized") == 4
    assert render.call_args_list[1].args[1] == StreamInfo(
        mimetype="image/svg+xml", extension=".svg", filename="vector image.SVG"
    )
    # The read-only table reader ignores drawings, including vector images.
    captured = []

    def capture(stream, info, **kwargs):
        captured.append(stream.read())
        return None

    _convert(_ImageConverter(capture), output.getvalue())
    assert captured == [_RED, svg, _BLUE, _BLUE]


def test_repeated_relationships_and_grouped_images_are_not_dropped() -> None:
    def group(name: str, content: bytes) -> bytes:
        if name != "xl/drawings/drawing1.xml":
            return content
        root = ET.fromstring(content)
        ns = _xlsx_images._NS
        anchor = root.find("xdr:oneCellAnchor", ns)
        assert anchor is not None
        picture = anchor.find("xdr:pic", ns)
        assert picture is not None
        anchor.remove(picture)
        group = ET.SubElement(anchor, "{" + ns["xdr"] + "}grpSp")
        nonvisual = ET.SubElement(group, "{" + ns["xdr"] + "}nvGrpSpPr")
        ET.SubElement(
            nonvisual, "{" + ns["xdr"] + "}cNvPr", {"id": "20", "name": "Group"}
        )
        ET.SubElement(nonvisual, "{" + ns["xdr"] + "}cNvGrpSpPr")
        ET.SubElement(group, "{" + ns["xdr"] + "}grpSpPr")
        group.append(picture)
        duplicate = ET.fromstring(ET.tostring(picture))
        properties = duplicate.find("xdr:nvPicPr/xdr:cNvPr", ns)
        assert properties is not None
        properties.set("id", "21")
        group.append(duplicate)
        return ET.tostring(root)

    captured = []

    def render(stream, info, **kwargs):
        captured.append(stream.read())
        return "<p>recognized</p>"

    result = _convert(_ImageConverter(render), _rewrite(_workbook(), group))
    assert captured == [_RED, _RED, _RED, _BLUE, _BLUE]
    assert "Image at " not in result
    assert result.count("recognized") == 5


def test_malformed_image_reference_does_not_become_silent_native_fallback() -> None:
    def missing_reference(name: str, content: bytes) -> bytes:
        if name != "xl/drawings/drawing1.xml":
            return content
        return content.replace(b'embed="rId1"', b'embed="missing"')

    with pytest.raises(KeyError):
        _convert(
            _ImageConverter(Mock(return_value="<p>recognized</p>")),
            _rewrite(_workbook(), missing_reference),
        )


def test_embedded_drawing_entities_are_rejected() -> None:
    def entity(name: str, content: bytes) -> bytes:
        if name != "xl/drawings/drawing1.xml":
            return content
        return b'<!DOCTYPE wsDr [<!ENTITY label "not expanded">]>' + content.replace(
            b'name="Image 1"', b'name="&label;"'
        )

    render = Mock(return_value="<p>recognized</p>")
    with pytest.raises(EntitiesForbidden):
        _convert(_ImageConverter(render), _rewrite(_workbook(), entity))
    render.assert_not_called()


def test_missing_dependencies_and_legacy_xls_stay_separate(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data = (Path(__file__).parent / "test_files" / "test.xls").read_bytes()
    legacy = XlsConverter()
    expected = legacy.convert(io.BytesIO(data), StreamInfo(extension=".xls")).markdown
    dependency = ImportError("openpyxl unavailable")
    monkeypatch.setattr(
        _xlsx_converter, "_xlsx_dependency_exc_info", (ImportError, dependency, None)
    )
    with pytest.raises(MissingDependencyException, match=r"\[xlsx\]"):
        XlsxConverter().convert(io.BytesIO(), _INFO)
    assert not hasattr(legacy, "_image_to_html")
    assert legacy.accepts(io.BytesIO(), StreamInfo(extension=".XLS"))
    assert (
        legacy.convert(io.BytesIO(data), StreamInfo(extension=".xls")).markdown
        == expected
    )


# Conversion regressions


def test_xlsx_legacy_show_zeroes_sheetview(tmp_path) -> None:
    from openpyxl import Workbook

    base_path = tmp_path / "base.xlsx"
    xlsx_path = tmp_path / "legacy_show_zeroes.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "hello"
    sheet["B1"] = "world"
    # A cell whose text is byte-identical to the malformed attribute. openpyxl stores it
    # as an inline string, so it lands in the very worksheet part that gets repaired.
    sheet["A2"] = ' showZeroes="0" '
    workbook.save(base_path)

    with zipfile.ZipFile(base_path) as source:
        with zipfile.ZipFile(xlsx_path, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(
                        b"<sheetView ", b'<sheetView showZeroes="0" ', 1
                    )
                    # Guard the fixture: the sheet view attribute and the cell text must
                    # both live here, or this test stops exercising the repair.
                    assert data.count(b'showZeroes="0"') == 2
                target.writestr(item, data)

    result = MarkItDown().convert(str(xlsx_path))

    assert "## Data" in result.markdown
    assert "hello" in result.markdown
    assert "world" in result.markdown

    # The repair must not reach into the worksheet's data: the cell keeps its text ...
    assert 'showZeroes="0"' in result.markdown
    # ... and no part of the document was silently renamed on the way through.
    assert "showZeros" not in result.markdown


def test_xlsx_show_zeroes_rename_is_scoped_to_sheet_view_tags() -> None:
    from markitdown.converters._xlsx_converter import _rename_show_zeroes_attribute

    worksheet_xml = (
        b"<worksheet><sheetViews>"
        b'<sheetView showZeroes="0" workbookViewId="0"/>'
        b'<sheetView\n\tshowZeroes="1"\ttabSelected="1"/>'
        b"</sheetViews>"
        # openpyxl ignores custom sheet views entirely, so they need no repair
        b'<customSheetViews><customSheetView showZeroes="0"/></customSheetViews>'
        b'<sheetData><row r="1">'
        b'<c r="A1" t="inlineStr"><is><t xml:space="preserve"> showZeroes="0" '
        b"</t></is></c>"
        b'<c r="B1"><f>IF(A1=" showZeroes=","x","y")</f><v>y</v></c>'
        b"</row></sheetData></worksheet>"
    )

    repaired = _rename_show_zeroes_attribute(worksheet_xml)

    # Every <sheetView> start tag is repaired, whatever separates its attributes ...
    assert repaired.count(b"showZeros=") == 2
    assert b'<sheetView showZeros="0" workbookViewId="0"/>' in repaired
    assert b'<sheetView\n\tshowZeros="1"\ttabSelected="1"/>' in repaired

    # ... and nothing outside those start tags is rewritten.
    assert b'<t xml:space="preserve"> showZeroes="0" </t>' in repaired
    assert b'<f>IF(A1=" showZeroes=","x","y")</f>' in repaired
    assert b'<customSheetView showZeroes="0"/>' in repaired


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__]))
