"""Rebuild missing-text PPTX fixtures from copies of test.pptx.

Requires python-pptx. The source fixture is never overwritten.
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


FIXTURES = Path(__file__).resolve().parent


def _missing_text_presentation() -> Presentation:
    # Modify a copy in memory; keep the checked-in deck unchanged.
    presentation = Presentation(FIXTURES / "test.pptx")
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Preserved slide title"
    slide.placeholders[1].text = "Preserved body text"
    slide.shapes.add_textbox(
        Inches(1), Inches(5), Inches(5), Inches(1)
    ).text = "Text after empty shape"
    slide.notes_slide.notes_text_frame.text = "Preserved speaker notes"
    return presentation


def main() -> None:
    for empty_shape in ("title", "body"):
        presentation = _missing_text_presentation()
        slide = presentation.slides[-1]
        shape = slide.shapes.title if empty_shape == "title" else slide.placeholders[1]
        shape.text_frame.clear()
        shape.text_frame.paragraphs[0].add_run().text = ""
        presentation.save(FIXTURES / f"pptx_empty_{empty_shape}_run.pptx")

    for notes_state in ("empty_run", "missing_body_frame"):
        presentation = _missing_text_presentation()
        notes_slide = presentation.slides[-1].notes_slide
        if notes_state == "empty_run":
            frame = notes_slide.notes_text_frame
            frame.clear()
            frame.paragraphs[0].add_run().text = ""
        else:
            body = notes_slide.notes_placeholder
            body._element.getparent().remove(body._element)
        presentation.save(FIXTURES / f"pptx_notes_{notes_state}.pptx")


if __name__ == "__main__":
    main()
